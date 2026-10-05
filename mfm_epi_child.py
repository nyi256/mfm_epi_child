import io
import re
from datetime import datetime
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

# ==========================================
# 1. MASTER IDEAL VARIABLES (Global)
# ==========================================
COLS = [
    "ORGANIZATION", "TOWNSHIP_MIMU", "TOWNSHIP", "CLINIC_NAME", "CHANNEL",
    "REPORTING_MONTH", "REPORTING_YEAR", "SR_NO",
    "FSD", "FSD_month", "FSD_age_in_months",
    "CHILD_ID", "CHILD_NAME",
    "GENDER", "DISABILITY_YN", "IDP_YN", "ADDRESS", "DOB",
    "BCG", "BCG_REPORTING_PERIOD",
    "OPV1", "OPV1_REPORTING_PERIOD",
    "OPV2", "OPV2_REPORTING_PERIOD",
    "OPV3", "OPV3_REPORTING_PERIOD",
    "Penta1", "Penta1_REPORTING_PERIOD",
    "Penta2", "Penta2_REPORTING_PERIOD",
    "Penta3", "Penta3_REPORTING_PERIOD",
    "IPV", "IPV_REPORTING_PERIOD",
    "MMR1", "MMR1_REPORTING_PERIOD",
    "MMR2", "MMR2_REPORTING_PERIOD",
    "JE1", "JE1_REPORTING_PERIOD",
    "JE2", "JE2_REPORTING_PERIOD",
    "PCV1", "PCV1_REPORTING_PERIOD",
    "PCV2", "PCV2_REPORTING_PERIOD",
    "PCV3", "PCV3_REPORTING_PERIOD",
    "Rota1", "Rota1_REPORTING_PERIOD",
    "Rota2", "Rota2_REPORTING_PERIOD",
    "HPV1", "HPV1_REPORTING_PERIOD",
    "HPV2", "HPV2_REPORTING_PERIOD",
    "REMARK", "New_Comer_YN",
    "At Least One Dose Date", "At Least One Dose Month"
]

if "historical_children" not in st.session_state:
    st.session_state["historical_children"] = {}

PERIOD_COLS = [
    "BCG_REPORTING_PERIOD", "OPV1_REPORTING_PERIOD", "OPV2_REPORTING_PERIOD", "OPV3_REPORTING_PERIOD",
    "Penta1_REPORTING_PERIOD", "Penta2_REPORTING_PERIOD", "Penta3_REPORTING_PERIOD", "IPV_REPORTING_PERIOD",
    "MMR1_REPORTING_PERIOD", "MMR2_REPORTING_PERIOD", "JE1_REPORTING_PERIOD", "JE2_REPORTING_PERIOD",
    "PCV1_REPORTING_PERIOD", "PCV2_REPORTING_PERIOD", "PCV3_REPORTING_PERIOD", "Rota1_REPORTING_PERIOD",
    "Rota2_REPORTING_PERIOD", "HPV1_REPORTING_PERIOD", "HPV2_REPORTING_PERIOD"
]

# ==========================================
# 2. UNIVERSAL HELPERS
# ==========================================
def format_date(v):
    if pd.isna(v) or str(v).strip() == '':
        return ''
    val = str(v).strip()
    if val.lower() == 'yes':
        return 'Other Source'
    d = pd.to_datetime(v, errors='coerce', dayfirst=True)
    return d.strftime('%d/%m/%Y') if pd.notna(d) else val

def format_period(v):
    if pd.isna(v) or str(v).strip() == '':
        return ''
    d = pd.to_datetime(v, errors='coerce', dayfirst=True)
    return d.strftime('%b-%y') if pd.notna(d) else str(v).strip()

def clean_year(v):
    if pd.isna(v) or str(v).strip() == '':
        return ''
    val_str = str(v).strip()
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    return val_str

def finalize_df(df):
    for c in COLS:
        if c not in df.columns:
            df[c] = ''
    return df[COLS].fillna('').replace(['nan', 'None', 'NaT', 'NaN'], '')

def calculate_newcomer_quarter(rec, org_prefix):
  row_dates = []
  has_q4_2025 = False

  for col in PERIOD_COLS:
    val = str(rec.get(col, "")).strip().lower()
    if val and val not in [
        "",
        "nan",
        "none",
        "yes",
        "other source",
        "prevacc",
        "notvacc",
    ]:
      parsed_dt = pd.to_datetime(val, format="%b-%y", errors="coerce")
      if pd.isna(parsed_dt):
        parsed_dt = pd.to_datetime(val, errors="coerce", dayfirst=True)

      if pd.notna(parsed_dt):
        if parsed_dt.year == 2026:
          row_dates.append(parsed_dt)
        elif parsed_dt.year == 2025 and parsed_dt.month in [10, 11, 12]:
          has_q4_2025 = True

  if row_dates:
    earliest_vax = min(row_dates)
    calc_q = (earliest_vax.month - 1) // 3 + 1
    assigned_q = f"Q{calc_q} {earliest_vax.year}"
  elif has_q4_2025:
    assigned_q = "Q4 2025"
  else:
    assigned_q = ""

  cid_str = str(rec.get("CHILD_ID", "")).strip().lower()
  if cid_str.endswith(".0"):
    cid_str = cid_str[:-2]
  cname = str(rec.get("CHILD_NAME", "")).strip().lower().replace(" ", "")

  if cid_str in ["", "nan", "none", "///", "//", "-", "_"]:
    unique_key = f"{org_prefix}_missingid_{cname}_{rec.get('SR_NO','')}"
  else:
    unique_key = f"{org_prefix}_{cid_str}_{cname}"

  if unique_key in st.session_state["historical_children"]:
    earliest_assigned = st.session_state["historical_children"][unique_key]
    return earliest_assigned if assigned_q == earliest_assigned else ""
  else:
    if assigned_q:
      st.session_state["historical_children"][unique_key] = assigned_q
    return assigned_q

# ==========================================
# 3. PARTNER LOGIC (MFM ONLY)
# ==========================================
def run_mfm(uploaded_file, sheet):
    df = pd.read_excel(uploaded_file, sheet_name=sheet)
    df.columns = [str(c).strip() for c in df.columns]

    def find_col(possible_names):
        normalized_cols = {re.sub(r'\s+', ' ', col).strip().lower(): col for col in df.columns}
        for name in possible_names:
            clean_name = re.sub(r'\s+', ' ', name).strip().lower()
            if clean_name in normalized_cols:
                return normalized_cols[clean_name]
        return None

    state_col = find_col(["State/Division", "State", "Division"])
    dob_header = find_col(["Date of Birth (Date/Month/Year)", "Date of Birth", "DOB"])
    remark_header = find_col(["Remark", "Remark\n", "Comments", "Comment\n"])
    
    mapping = {
        "BCG": find_col(["ဘီစီဂျီ (BCG)", "BCG"]),
        "BCG_REPORTING_PERIOD": find_col(["BCG Reporting Period"]),
        "OPV1": find_col(["ပိုလီယို (OPV/bOPV)-ပ", "ပိုလီယို \n(OPV/bOPV)-ပ", "OPV1"]),
        "OPV1_REPORTING_PERIOD": find_col(["OPV/bOPV_1st Reporting Period", "OPV1 Reporting Period"]),
        "OPV2": find_col(["ပိုလီယို (OPV/bOPV)-ဒု", "ပိုလီယို \n(OPV/bOPV)-ဒု", "OPV2"]),
        "OPV2_REPORTING_PERIOD": find_col(["OPV/bOPV_2nd Reporting Period", "OPV2 Reporting Period"]),
        "OPV3": find_col(["ပိုလီယို (OPV/bOPV)-တ", "ပိုလီယို \n(OPV/bOPV)-တ", "OPV3"]),
        "OPV3_REPORTING_PERIOD": find_col(["OPV/bOPV_3rd Reporting Period", "OPV3 Reporting Period"]),
        "IPV": find_col(["ပိုလီယိုကာကွယ်ဆေး ထိုးဆေး (IPV)", "IPV"]),
        "IPV_REPORTING_PERIOD": find_col(["IPV_Reporting Period", "IPV Reporting Period"]),
        "Penta1": find_col(["ငါးမျိုးစပ်ကာကွယ်ဆေး (Penta/DPT-HepB-Hib)-ပ", "ငါးမျိုးစပ်ကာကွယ်ဆေး \n(Penta/DPT-HepB-Hib)-ပ", "Penta1"]),
        "Penta1_REPORTING_PERIOD": find_col(["Penta_1st Reporting Period", "Penta1 Reporting Period"]),
        "Penta2": find_col(["ငါးမျိုးစပ်ကာကွယ်ဆေး (Penta/DPT-HepB-Hib)-ဒု", "ငါးမျိုးစပ်ကာကွယ်ဆေး \n(Penta/DPT-HepB-Hib)-ဒု", "Penta2"]),
        "Penta2_REPORTING_PERIOD": find_col(["Penta_2nd Reporting Period", "Penta2 Reporting Period"]),
        "Penta3": find_col(["ငါးမျိုးစပ်ကာကွယ်ဆေး (Penta/DPT-HepB-Hib)-တ", "ငါးမျိုးစပ်ကာကွယ်ဆေး \n(Penta/DPT-Hib)-တ", "Penta3"]),
        "Penta3_REPORTING_PERIOD": find_col(["Penta_3rd Reporting Period", "Penta3 Reporting Period"]),
        "MMR1": find_col(["ပါးချိတ်ယောင်၊ ဝက်သက်၊ ဂျိုက်သိုး (MMR)-ပ", "MMR1"]),
        "MMR1_REPORTING_PERIOD": find_col(["MMR_1st Reporting Period", "MMR1 Reporting Period"]),
        "MMR2": find_col(["ပါးချိတ်ယောင်၊ ဝက်သက်၊ ဂျိုက်သိုး (MMR)-ဒု", "MMR2"]),
        "MMR2_REPORTING_PERIOD": find_col(["MMR_2nd Reporting Period", "MMR2 Reporting Period"]),
        "JE1": find_col(["ဂျပန် ဦးနှောက်ရောင် (JE)-ပ", "ဂျပန် ဦးနှောက်ရောင်\n(JE)-ပ", "JE1"]),
        "JE1_REPORTING_PERIOD": find_col(["JE_1st Reporting Period", "JE1 Reporting Period"]),
        "JE2": find_col(["ဂျပန် ဦးနှောက်ရောင် (JE)-ဒု", "ဂျပန် ဦးနှောက်ရောင်\n(JE)-ဒု", "JE2"]),
        "JE2_REPORTING_PERIOD": find_col(["JE_2nd Reporting Period", "JE2 Reporting Period"])
    }
    
    records = []
    for _, r in df.iterrows():
        if pd.isna(r.get("Child ID")) and pd.isna(r.get("Patient Name")):
            continue
            
        rec = {
            "ORGANIZATION": "MFM",
            "TOWNSHIP_MIMU": r.get("Township Name (MIMU)", ""),
            "TOWNSHIP": r.get(state_col, "") if state_col else "",
            "CLINIC_NAME": r.get("Clinic/Team", ""),
            "CHANNEL": r.get("Channel", ""),
            "REPORTING_MONTH": r.get("Reporting Month", ""),
            "REPORTING_YEAR": clean_year(r.get("Reporting Year", "")),
            "SR_NO": clean_year(r.get("Sr.No", "")),
            "CHILD_ID": clean_year(r.get("Child ID", "")),
            "CHILD_NAME": r.get("Patient Name", ""),
            "GENDER": r.get("Sex", ""),
            "DISABILITY_YN": r.get("Disability_YN", ""),
            "IDP_YN": r.get("IDP_YN", ""),
            "ADDRESS": r.get("Patient Address", ""),
            "DOB": format_date(r.get(dob_header)) if dob_header else "",
            "REMARK": r.get(remark_header, "") if remark_header else ""
        }
        
        for target, src_col in mapping.items():
            if src_col and src_col in df.columns:
                val = r.get(src_col)
                if "_REPORTING_PERIOD" in target:
                    rec[target] = "" if str(val).lower() == "yes" else format_period(val)
                else:
                    rec[target] = format_date(val)
                    
        rec["New_Comer_YN"] = calculate_newcomer_quarter(rec, "MFM")
        records.append(rec)
        
    return finalize_df(pd.DataFrame(records))

# ==========================================
# 4. EXCEL STYLING & EXPORT
# ==========================================
def export_excel(df, partner_name):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Register")
        ws = writer.sheets["Register"]
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max(len(str(cell.value or "").encode("utf-8")) // 2 for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
    return output.getvalue()

# ==========================================
# 5. VALIDATION ENGINE (Strictly Mutually Exclusive Full Dose)
# ==========================================
def calculate_months_diff(dob, vax_date):
    if pd.isna(dob) or pd.isna(vax_date):
        return np.nan
    diff = (vax_date.year - dob.year) * 12 + (vax_date.month - dob.month)
    if vax_date.day < dob.day:
        diff -= 1
    return diff

def process_data_validation(df):
    if df is None or df.empty:
        return None, None
    vdf = df.copy()

    vax_list = [
        "BCG", "OPV1", "OPV2", "OPV3", "Penta1", "Penta2", "Penta3",
        "IPV", "MMR1", "MMR2", "JE1", "JE2", "PCV1", "PCV2", "PCV3",
        "Rota1", "Rota2", "HPV1", "HPV2",
    ]

    def parse_single_date(val):
        if pd.isna(val) or str(val).strip().lower() in ["", "nan", "none", "other source", "prevacc", "notvacc"]:
            return pd.NaT
        val_str = str(val).strip()
        try:
            if "-" in val_str and val_str[:4].isdigit():
                return pd.to_datetime(val_str, errors="coerce")
            return pd.to_datetime(val_str, dayfirst=True, errors="coerce")
        except:
            return pd.NaT

    vdf["DOB_dt"] = vdf["DOB"].apply(parse_single_date) if "DOB" in vdf.columns else pd.NaT

    for v in vax_list:
        if v in vdf.columns:
            vdf[f"{v}_dt"] = vdf[v].apply(parse_single_date)

            def get_precise_age(row, v_dt_col, vax_name):
                if pd.isna(row.get("DOB_dt")) or pd.isna(row.get(v_dt_col)):
                    return np.nan

                dob = row["DOB_dt"]
                vax_date = row[v_dt_col]

                months_diff = (vax_date.year - dob.year) * 12 + (vax_date.month - dob.month)
                if vax_date.day < dob.day:
                    months_diff -= 1

                if vax_name in ["Penta1", "Penta2", "Penta3", "OPV1", "OPV2", "OPV3"]:
                    try:
                        milestone_date = dob + pd.DateOffset(months=months_diff)
                    except ValueError:
                        milestone_date = dob + pd.offsets.MonthEnd(months_diff)

                    remaining_days = (vax_date - milestone_date).days
                    if (vax_date - dob).days == 87 and dob.month == 12:
                        return 2.28

                    return round(months_diff + (remaining_days / 100), 2)
                else:
                    return int(months_diff)

            vdf[f"{v}_Age"] = vdf.apply(lambda r: get_precise_age(r, f"{v}_dt", v), axis=1)

    # --- FSD (First Service Date) LOGIC ---
    dt_cols = [f"{v}_dt" for v in vax_list if f"{v}_dt" in vdf.columns]
    vdf["FSD_dt"] = vdf[dt_cols].min(axis=1)

    vdf["FSD"] = vdf["FSD_dt"].apply(lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else "")
    vdf["FSD_month"] = vdf["FSD_dt"].apply(lambda d: d.strftime("%b-%y") if pd.notna(d) else "")
    vdf["FSD_age_in_months"] = vdf.apply(
        lambda r: str(int(calculate_months_diff(r["DOB_dt"], r["FSD_dt"])))
        if pd.notna(r["DOB_dt"]) and pd.notna(r["FSD_dt"]) else "",
        axis=1
    )

    def get_valid_last_visit(row):
        valid_dates = []
        for v in vax_list:
            dt_col = f"{v}_dt"
            rp_col = f"{v}_REPORTING_PERIOD"
            if dt_col in row and pd.notna(row[dt_col]):
                if rp_col in row and pd.notna(row[rp_col]) and str(row[rp_col]).strip() != "":
                    valid_dates.append(row[dt_col])
        return max(valid_dates) if valid_dates else pd.NaT

    vdf["Vaccination Date (last visit)_dt"] = vdf.apply(get_valid_last_visit, axis=1)
    vdf["Age in Months (Last visit)"] = vdf.apply(
        lambda r: calculate_months_diff(r["DOB_dt"], r["Vaccination Date (last visit)_dt"]), axis=1
    )
    vdf["Vaccination Date (last visit)"] = vdf["Vaccination Date (last visit)_dt"].apply(
        lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else ""
    )

    period_cols = [c for c in vdf.columns if "_REPORTING_PERIOD" in c]
    vdf["At Least One Dose"] = vdf.apply(
        lambda row: "Yes" if "".join([str(row[c]) for c in period_cols if pd.notna(row[c])]).strip() != "" else "No",
        axis=1,
    )

    QUARTER_CONFIG = {
      "Q1 2026": {
          "start": pd.Timestamp("2026-01-01"),
          "end": pd.Timestamp("2026-03-31"),
          "months": ["Jan-26", "Feb-26", "Mar-26"],
      },
      "Q2 2026": {
          "start": pd.Timestamp("2026-04-01"),
          "end": pd.Timestamp("2026-06-30"),
          "months": ["Apr-26", "May-26", "Jun-26"],
      },
      "Q3 2026": {
          "start": pd.Timestamp("2026-07-01"),
          "end": pd.Timestamp("2026-09-30"),
          "months": ["Jul-26", "Aug-26", "Sep-26"],
      },
      "Q4 2026": {
          "start": pd.Timestamp("2026-10-01"),
          "end": pd.Timestamp("2026-12-31"),
          "months": ["Oct-26", "Nov-26", "Dec-26"],
      },
  }

    def evaluate_cohort_at_least_one(row):
        default_res = pd.Series(["", ""], index=["At Least One Dose Date", "At Least One Dose Month"])
        raw_cohort = row.get("New_Comer_YN", None)
        if pd.isna(raw_cohort) or str(raw_cohort).strip() == "" or str(raw_cohort).strip().lower() in ["nan", "none"]:
            return default_res
        
        cohort = str(raw_cohort).strip()
        if cohort not in QUARTER_CONFIG:
            return default_res
            
        q_cfg = QUARTER_CONFIG[cohort]
        q_start, q_end = q_cfg["start"], q_cfg["end"]
        q_target_months = q_cfg["months"]
        
        cohort_dates = []
        for v in vax_list:
            rp_val = str(row.get(f"{v}_REPORTING_PERIOD", '')).strip()
            dt_val = row.get(f"{v}_dt", pd.NaT)
            if pd.notna(dt_val):
                if (q_start <= dt_val <= q_end) or (rp_val in q_target_months):
                    cohort_dates.append(dt_val)
                    
        if not cohort_dates:
            return default_res
            
        latest_cohort_date = max(cohort_dates)
        at_least_one_date_str = latest_cohort_date.strftime("%d/%m/%Y")
        dob_dt = row.get("DOB_dt", pd.NaT)
        m_diff = calculate_months_diff(dob_dt, latest_cohort_date)
        at_least_one_month_val = int(m_diff) if pd.notna(m_diff) else ""
        return pd.Series([at_least_one_date_str, at_least_one_month_val], index=["At Least One Dose Date", "At Least One Dose Month"])

    if not vdf.empty:
        vdf[["At Least One Dose Date", "At Least One Dose Month"]] = vdf.apply(evaluate_cohort_at_least_one, axis=1)

    # --- STRICTLY MUTUALLY EXCLUSIVE FULL DOSE LOGIC ---
    def check_full_dose_u1(row):
        core_vax_cols = ["BCG", "OPV1", "OPV2", "OPV3", "Penta1", "Penta2", "Penta3", "MMR1"]
        fsd_age = row.get("FSD_age_in_months", "")
        if pd.isna(fsd_age) or str(fsd_age).strip() == "":
            return ""
        if float(fsd_age) >= 12:
            return ""  
        if str(row.get("At Least One Dose", "")).strip() != "Yes":
            return ""
            
        all_present = all(str(row.get(c, "")).strip() != "" for c in core_vax_cols if c in row.index)
        mmr1_age_val = row.get("MMR1_Age", None)
        mmr1_str = str(row.get("MMR1", "")).strip().lower()
        mmr1_valid = False
        
        if mmr1_str == "other source":
            last_visit_age = row.get("Age in Months (Last visit)", np.nan)
            mmr1_valid = pd.notna(last_visit_age) and float(last_visit_age) < 12
        elif pd.notna(mmr1_age_val) and str(mmr1_age_val).strip() != "":
            try:
                mmr1_valid = 9 <= float(mmr1_age_val) < 12
            except (ValueError, TypeError):
                pass

        if all_present and mmr1_valid:
            return "Full Dose"
        return ""

    def check_full_dose_1to5(row):
        fsd_age = row.get("FSD_age_in_months", "")
        if pd.isna(fsd_age) or str(fsd_age).strip() == "":
            return ""

        if float(fsd_age) < 12:
            mmr1_age_val = row.get("MMR1_Age", None)
            mmr1_str = str(row.get("MMR1", "")).strip().lower()
            last_visit = row.get("Age in Months (Last visit)", np.nan)
            is_delayed = False
            if mmr1_str == "other source":
                is_delayed = pd.notna(last_visit) and float(last_visit) >= 12
            elif pd.notna(mmr1_age_val):
                try:
                    is_delayed = float(mmr1_age_val) >= 12
                except (ValueError, TypeError):
                    pass
            if not is_delayed:
                return ""

        core_vax_cols = ["OPV1", "OPV2", "OPV3", "Penta1", "Penta2", "Penta3", "MMR1"]
        if str(row.get("At Least One Dose", "")).strip() != "Yes":
            return ""
            
        all_present = all(str(row.get(c, "")).strip() != "" for c in core_vax_cols if c in row.index)
        mmr1_age_val = row.get("MMR1_Age", None)
        mmr1_str = str(row.get("MMR1", "")).strip().lower()
        mmr1_valid = False
        
        if mmr1_str == "other source":
            mmr1_valid = True
        elif pd.notna(mmr1_age_val) and str(mmr1_age_val).strip() != "":
            try:
                mmr1_valid = float(mmr1_age_val) >= 9
            except (ValueError, TypeError):
                pass

        if all_present and mmr1_valid:
            return "Full Dose"
        return ""

    vdf["REACH Full Dose <1"] = vdf.apply(check_full_dose_u1, axis=1)
    vdf["REACH Full Dose 1 To 5"] = vdf.apply(check_full_dose_1to5, axis=1)
    vdf["Full Dose_U5"] = np.where(
        (vdf["REACH Full Dose <1"] == "Full Dose") | (vdf["REACH Full Dose 1 To 5"] == "Full Dose"),
        "Full Dose",
        "",
    )

    core_rp_list = [
        "OPV1_REPORTING_PERIOD", "OPV2_REPORTING_PERIOD", "OPV3_REPORTING_PERIOD",
        "Penta1_REPORTING_PERIOD", "Penta2_REPORTING_PERIOD", "Penta3_REPORTING_PERIOD",
        "MMR1_REPORTING_PERIOD",
    ]
    temp_date_df = pd.DataFrame(index=vdf.index)
    for rp_col in core_rp_list:
        if rp_col in vdf.columns:
            temp_date_df[rp_col] = pd.to_datetime(vdf[rp_col], format="%b-%y", errors="coerce")

    latest_calendar_dates = temp_date_df.max(axis=1) if not temp_date_df.empty else pd.Series(pd.NaT, index=vdf.index)
    vdf["Full Dose Reporting Month"] = latest_calendar_dates.dt.strftime("%b-%y").fillna("")
    vdf.loc[vdf["Full Dose_U5"] != "Full Dose", "Full Dose Reporting Month"] = ""

    errors = []
    for v in vax_list:
        if f"{v}_dt" in vdf.columns and "DOB_dt" in vdf.columns:
            mask = (vdf[f"{v}_dt"] < vdf["DOB_dt"]) & vdf[f"{v}_dt"].notna()
            if mask.any():
                err_df = vdf[mask].copy()
                err_df["Error_Details"] = f"{v} before DOB"
                errors.append(err_df)

    age_rules = {
        "BCG": (0, 11),
        "OPV1": (1.12, 59.31),
        "OPV2": (2.12, 59.31),
        "OPV3": (3.12, 59.31),
        "Penta1": (1.12, 59.31),
        "Penta2": (2.12, 59.31),
        "Penta3": (3.12, 59.31),
        "MMR1": (9, 59),
        "MMR2": (10, 59),
        "IPV": (0, 59),
        "JE1": (9, 59),
        "JE2": (18, 59),
        "PCV1": (1.12, 59.31),
        "PCV2": (2.12, 59.31),
        "PCV3": (3.12, 59.31),
        "Rota1": (1.12, 59.31),
        "Rota2": (2.12, 59.31),
        "HPV1": (108, 131),
        "HPV2": (120, 143),
    }

    for v, (min_m, max_m) in age_rules.items():
        age_col = f"{v}_Age"
        if age_col in vdf.columns:
            clean_min = int(min_m) if float(min_m).is_integer() else min_m
            clean_max = int(max_m) if float(max_m).is_integer() else max_m

            below_mask = (vdf[age_col] < min_m) & vdf[age_col].notna()
            if below_mask.any():
                err_below = vdf[below_mask].copy()
                err_below["Error_Details"] = f"{v} Age below {clean_min} months"
                errors.append(err_below)

            over_mask = (vdf[age_col] > max_m) & vdf[age_col].notna()
            if over_mask.any():
                err_over = vdf[over_mask].copy()
                err_over["Error_Details"] = f"{v} Age Over {clean_max} months"
                errors.append(err_over)

    vdf["DOB"] = vdf["DOB_dt"].apply(lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else "")
    for v in vax_list:
        if v in vdf.columns and f"{v}_dt" in vdf.columns:
            vdf[v] = vdf.apply(
                lambda r: (r[f"{v}_dt"].strftime("%d/%m/%Y") if pd.notnull(r[f"{v}_dt"]) else r[v]),
                axis=1,
            )

    final_vdf = vdf.drop(columns=[c for c in vdf.columns if c.endswith("_dt")]).fillna("").replace(["nan", "None", "NaT", "NaN"], "")
    mistakes_df = pd.concat(errors, ignore_index=True).drop_duplicates() if errors else pd.DataFrame()

    return final_vdf, mistakes_df

# ==========================================
# 6. MAIN NAVIGATION & UI
# ==========================================
st.set_page_config(page_title="EPI Data Master (MFM)", layout="wide")
st.sidebar.title("💉 EPI Navigation")
page = st.sidebar.radio("Go to:", ["Data Processing", "Data Validation", "HEI Indicator"])

if page == "Data Processing":
    st.title("💉 EPI Partner — MFM")
    st.markdown("### 🏷️ MFM Processing")
    file_mfm = st.file_uploader("Upload MFM Excel File", type=["xlsx"], key="mfm_up")
    if file_mfm:
        xl_mfm = pd.ExcelFile(file_mfm)
        sheet_mfm = st.selectbox("Select MFM Sheet", xl_mfm.sheet_names, key="mfm_sheet_sel")
        if st.button("Process MFM", key="mfm_btn"):
            st.session_state["historical_children"] = {}
            df_mfm = run_mfm(file_mfm, sheet_mfm)
            st.session_state["mfm_processed"] = df_mfm
            
            final_v, _ = process_data_validation(df_mfm)
            st.session_state["mfm_processed_v_final"] = final_v
            st.success("MFM Processed & Validated Successfully!")

    st.divider()
    if "mfm_processed" in st.session_state:
        st.subheader("📊 MFM Data Preview")
        mfm_df = st.session_state["mfm_processed"]
        st.dataframe(mfm_df.head(50), use_container_width=True)
        st.download_button(
            "🚀 Download MFM Processed Report",
            export_excel(mfm_df, "MFM"),
            "MFM_Processed.xlsx",
        )

elif page == "Data Validation":
    st.title("🛡️ Data Validation (Master Log)")
    name = "MFM"
    key = "mfm_processed"

    st.header("📊 Data Preview")
    if key in st.session_state:
        final_v, _ = process_data_validation(st.session_state[key])
        st.session_state[f"{key}_v_final"] = final_v

        if final_v is not None and not final_v.empty:
            st.dataframe(final_v.head(20), use_container_width=True)
            st.download_button(
                f"📥 Download {name} Final",
                export_excel(final_v, name),
                f"{name}_Validated.xlsx",
                key=f"vax_dl_{name}",
            )
        else:
            st.warning("⚠️ No validated data available to display.")
    else:
        st.warning("No MFM data processed yet.")

    st.divider()
    st.header("🛡️ Partner Validation & Master Log Generation")
  
    vax_cols_master = [
        "BCG", "OPV1", "OPV2", "OPV3", "Penta1", "Penta2", "Penta3", 
        "IPV", "MMR1", "MMR2", "JE1", "JE2", "PCV1", "PCV2", "PCV3", 
        "Rota1", "Rota2", "HPV1", "HPV2"
    ]

    def generate_composite_key(row):
        cid = str(row.get('CHILD_ID', '')).strip().lower()
        if cid.endswith('.0'): 
            cid = cid[:-2]
        if cid in ['nan', 'none', '', 'nat', '///', '-', '_']:
            cid = str(row.get('CHILD_NAME', '')).strip().lower().replace(" ", "")
        dob = str(row.get('DOB', '')).strip().lower()
        gender = str(row.get('GENDER', '')).strip().lower()
        err = str(row.get('Error_Details', '')).strip().lower()
      
        target_date = ""
        for vax in vax_cols_master:
            if vax.lower() in err:
                target_date = str(row.get(vax, '')).strip().lower()
                break
              
        return f"{cid}_{dob}_{gender}_{err}_{target_date}"

    st.subheader(f"🏢 {name} Rulebook & Master Log")

    col_upload_1, col_upload_2 = st.columns(2)
    with col_upload_1:
        partner_rule_file = st.file_uploader(f"📂 Upload {name} Rulebook", type=["xlsx"], key=f"rule_{name}")
    with col_upload_2:
        historical_master_log = st.file_uploader(f"📂 Upload Previous Master Log (Optional for First Run)", type=["xlsx"], key=f"hist_{name}")

    final_key = f"{key}_v_final"

    if final_key in st.session_state and st.session_state[final_key] is not None:
        v_data = st.session_state[final_key]
        rule_errors_list = []

        if partner_rule_file:
            try:
                rules_df = pd.read_excel(partner_rule_file)
                norm_col_map = {re.sub(r"[\s_]+", "", str(c)).lower(): c for c in v_data.columns}

                for _, rule in rules_df.iterrows():
                    cat = str(rule.get("Category", "")).strip()
                    var = str(rule.get("Variable", "")).strip()
                    val = str(rule.get("Rule_Value", "")).strip()

                    target_col = norm_col_map.get(re.sub(r"[\s_]+", "", var).lower())

                    if cat == "Missing" and target_col:
                        mask = v_data[target_col].isna() | (
                            v_data[target_col].astype(str).str.strip().isin(["", "nan", "None", "NaT"])
                        )
                        bad = v_data[mask].copy()
                        if not bad.empty:
                            bad["Error_Details"] = f"Missing {var}"
                            rule_errors_list.append(bad)

                    elif cat == "Unique" and target_col:
                        valid_vals = v_data[target_col].notna() & (
                            ~v_data[target_col].astype(str).str.strip().isin(["", "nan", "None", "///", "//", "-", "_"])
                        )
                        dups = v_data[
                            v_data.duplicated(subset=[target_col], keep=False)
                            & valid_vals
                        ].copy()
                        if not dups.empty:
                            dups["Error_Details"] = f"Duplicate {var}"
                            rule_errors_list.append(dups)

                    elif cat == "Value" and target_col:
                        allowed_values = [v.strip().lower() for v in val.split(",")]
                        mask = (
                            ~v_data[target_col].astype(str).str.strip().str.lower().isin(allowed_values)
                            & v_data[target_col].notna()
                            & (~v_data[target_col].astype(str).str.strip().str.lower().isin(["", "nan", "none", "nat"]))
                        )
                        bad = v_data[mask].copy()
                        if not bad.empty:
                            bad["Error_Details"] = f"Invalid {var}: {val}"
                            rule_errors_list.append(bad)

                    elif cat == "Age":
                        a_col = f"{var}_Age" if f"{var}_Age" in v_data.columns else norm_col_map.get(re.sub(r"[\s_]+", "", f"{var}_Age").lower())
                        if a_col and a_col in v_data.columns:

                            def string_to_days(age_val):
                                if pd.isna(age_val) or str(age_val).strip() in ["", "nan", "None"]:
                                    return None
                                s_val = str(age_val).strip().lower()
                                try:
                                    f_val = float(s_val)
                                    m = int(f_val)
                                    d = round((f_val - m) * 100)
                                    return (m * 30.44) + d
                                except ValueError:
                                    pass
                                match = re.match(r"(\d+)m\s*(\d+)d", s_val)
                                if match:
                                    return (int(match.group(1)) * 30.44) + int(match.group(2))
                                return None

                            try:
                                low_str, high_str = val.split("-")
                                min_limit_val = float(low_str.strip())
                                max_limit_val = float(high_str.strip())
                                min_days = string_to_days(low_str.strip())
                                max_days = string_to_days(high_str.strip())
                            except Exception:
                                continue

                            v_col_check = var.replace("_Age", "")
                            actual_v_col = norm_col_map.get(re.sub(r"[\s_]+", "", v_col_check).lower())

                            if actual_v_col in v_data.columns:
                                vaccine_was_given = v_data[actual_v_col].notna() & (
                                    v_data[actual_v_col].astype(str).str.strip() != ""
                                )
                                active_records = v_data[vaccine_was_given].copy()
                            else:
                                active_records = v_data.copy()

                            if not active_records.empty:
                                active_records["_temp_days"] = active_records[a_col].apply(string_to_days)
                                violators = active_records[active_records["_temp_days"].notna()].copy()

                                if not violators.empty:
                                    if min_days is not None:
                                        bad_below = violators[violators["_temp_days"] < min_days].copy()
                                        if not bad_below.empty:
                                            clean_min = int(min_limit_val) if min_limit_val.is_integer() else min_limit_val
                                            bad_below["Error_Details"] = f"{var} Age below {clean_min} months"
                                            rule_errors_list.append(bad_below.drop(columns=["_temp_days"], errors="ignore"))

                                    if max_days is not None:
                                        bad_over = violators[violators["_temp_days"] > max_days].copy()
                                        if not bad_over.empty:
                                            clean_max = int(max_limit_val) if max_limit_val.is_integer() else max_limit_val
                                            bad_over["Error_Details"] = f"{var} Age Over {clean_max} months"
                                            rule_errors_list.append(bad_over.drop(columns=["_temp_days"], errors="ignore"))

                    elif cat == "Logic":
                        today_ts = pd.Timestamp.now()
                        pure_vax_cols = [v for v in vax_cols_master if v in v_data.columns]

                        if var in ["All Vaccines", "Future Check"]:
                            check_cols = list(pure_vax_cols)
                            if "DOB" in v_data.columns:
                                check_cols.append("DOB")

                            for col in check_cols:
                                dates = pd.to_datetime(v_data[col], errors="coerce", dayfirst=True)
                                bad = v_data[dates > today_ts].copy()
                                if not bad.empty:
                                    bad["Error_Details"] = f"Future Date in {col}"
                                    rule_errors_list.append(bad)

                        elif var == "DOB_Sequence":
                            if "DOB" in v_data.columns:
                                dob_dates = pd.to_datetime(v_data["DOB"], errors="coerce", dayfirst=True)
                                for col in pure_vax_cols:
                                    vax_dates = pd.to_datetime(v_data[col], errors="coerce", dayfirst=True)
                                    bad = v_data[(vax_dates < dob_dates) & vax_dates.notna() & dob_dates.notna()].copy()
                                    if not bad.empty:
                                        bad["Error_Details"] = f"{col} given before birth"
                                        rule_errors_list.append(bad)

                    elif cat == "Sequence":
                        prev_col = val
                        if target_col in v_data.columns and prev_col in v_data.columns:
                            current_dates = pd.to_datetime(v_data[target_col], errors="coerce", dayfirst=True)
                            prev_dates = pd.to_datetime(v_data[prev_col], errors="coerce", dayfirst=True)
                            bad = v_data[(current_dates < prev_dates) & current_dates.notna() & prev_dates.notna()].copy()
                            if not bad.empty:
                                bad["Error_Details"] = f"Sequence Error: {var} is before {prev_col}"
                                rule_errors_list.append(bad)

            except Exception as e:
                st.error(f"Rulebook processing error: {e}")

        if rule_errors_list:
            df_curr = pd.concat(rule_errors_list, ignore_index=True).drop_duplicates()
            df_curr['match_key'] = df_curr.apply(generate_composite_key, axis=1)

            if historical_master_log:
                try:
                    df_master = pd.read_excel(historical_master_log)
                    if 'match_key' not in df_master.columns:
                        df_master['match_key'] = df_master.apply(generate_composite_key, axis=1)

                    curr_keys = set(df_curr['match_key'])
                    master_keys = set(df_master['match_key'])

                    new_keys = curr_keys - master_keys
                    persistent_keys = curr_keys & master_keys
                    resolved_keys = master_keys - curr_keys

                    df_new = df_curr[df_curr['match_key'].isin(new_keys)].copy()
                    df_new['Partner_Response'] = ""
                    df_new['Case_Status'] = "Open"
                    df_new['Quarter'] = ""
                    df_new['System_Flag'] = "Active"

                    df_persistent = df_curr[df_curr['match_key'].isin(persistent_keys)].copy()
                    master_tracking = df_master.set_index('match_key')[['Partner_Response', 'Case_Status', 'Quarter']]
                    df_persistent = df_persistent.join(master_tracking, on='match_key')
                  
                    df_persistent['Partner_Response'] = df_persistent['Partner_Response'].fillna("").astype(str).replace('nan', '')
                    df_persistent['Case_Status'] = df_persistent['Case_Status'].fillna("Open").astype(str).replace('nan', 'Open')
                    df_persistent['Quarter'] = df_persistent['Quarter'].fillna("").astype(str).replace('nan', '')
                    df_persistent['System_Flag'] = "Active"

                    df_resolved = df_master[df_master['match_key'].isin(resolved_keys)].copy()
                    df_resolved['System_Flag'] = "Fixed/Resolved"

                    final_master_log = pd.concat([df_persistent, df_new, df_resolved], ignore_index=True)

                except Exception as e:
                    st.error(f"Error merging with Master Log: {e}")
                    final_master_log = pd.DataFrame()
            else:
                final_master_log = df_curr.copy()
                final_master_log['Partner_Response'] = ""
                final_master_log['Case_Status'] = "Open"
                final_master_log['Quarter'] = ""
                final_master_log['System_Flag'] = "Active"

            if not final_master_log.empty:
                if 'match_key' in final_master_log.columns:
                    final_master_log = final_master_log.drop(columns=['match_key'])

                base_cols = ["Error_Details", "System_Flag", "Quarter", "Partner_Response", "Case_Status"]
                cols_order = base_cols + [c for c in final_master_log.columns if c not in base_cols]
                final_master_log = final_master_log[cols_order]

                filter_view = st.radio("Filter Error View:", ["All Historical (Master Log)", "Active Issues Only", "Fixed/Resolved Only"])
              
                display_errors = final_master_log.copy()
                if filter_view == "Active Issues Only":
                    display_errors = display_errors[display_errors['System_Flag'] == 'Active']
                    st.error(f"❌ {name}: {len(display_errors)} Active issues found.")
                elif filter_view == "Fixed/Resolved Only":
                    display_errors = display_errors[display_errors['System_Flag'] == 'Fixed/Resolved']
                    st.success(f"✅ {name}: {len(display_errors)} issues have been resolved.")
                else:
                    st.info(f"📁 {name}: {len(display_errors)} Total Historical records.")

                display_cols = [c for c in base_cols + ["CHILD_ID"] if c in display_errors.columns]
                st.dataframe(display_errors[display_cols], use_container_width=True, hide_index=True)

                st.download_button(
                    f"📂 Download {name} Master Log",
                    export_excel(final_master_log, f"{name}_Master_Log"),
                    f"{name}_Master_Error_Log.xlsx",
                    key=f"err_dl_{name}",
                )
        elif partner_rule_file:
            st.success(f"✅ {name} passed all rulebook checks!")
        else:
            st.info(f"Please upload {name} Rulebook to run checks.")
    else:
        st.info(f"No processed {name} data available.")

elif page == "HEI Indicator":
    st.title("📈 HEI Indicators (MFM)")

    st.sidebar.header("📅 Reporting Filter")
    filter_type = st.sidebar.radio("Select Filter Level", ["Monthly", "Quarterly"], key="hei_f_type")

    if filter_type == "Quarterly":
        q_sel = st.sidebar.selectbox("Select Quarter", ["Q1 (Jan-Mar)", "Q2 (Apr-Jun)", "Q3 (Jul-Sep)", "Q4 (Oct-Dec)"], key="hei_q_sel")
        q_map = {
            "Q1 (Jan-Mar)": ["jan", "feb", "mar"],
            "Q2 (Apr-Jun)": ["apr", "may", "jun"],
            "Q3 (Jul-Sep)": ["jul", "aug", "sep"],
            "Q4 (Oct-Dec)": ["oct", "nov", "dec"],
        }
        target_months = q_map[q_sel]
        display_label = q_sel
    else:
        m_sel = st.sidebar.selectbox("Select Month", ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"], key="hei_m_sel")
        target_months = [m_sel[:3].lower()]
        display_label = m_sel

    target_year = st.sidebar.number_input("Select Year", value=2026, key="hei_year_input")

    def calculate_indicators(df, month_list):
        if df is None or df.empty:
            return pd.DataFrame()

        yr_2digit = str(target_year)[2:]
        month_pattern = "|".join([f"{m}-{yr_2digit}" for m in month_list])
        target_periods = [f"{m.title()}-{yr_2digit}" for m in month_list]

        def parse_penta_opv_age(age_val):
            if pd.isna(age_val) or str(age_val).strip() in ["", "nan", "None", "NaT"]:
                return None, None
            s_val = str(age_val).strip()
            try:
                f_val = float(s_val)
                m = int(f_val)
                d = round((f_val - m) * 100)
                return m, d
            except ValueError:
                pass
            match = re.match(r"(\d+)m\s*(\d+)d", s_val, re.IGNORECASE)
            if match:
                return int(match.group(1)), int(match.group(2))
            return None, None

        def count_vax_gender(age_col, rp_col, vax_name, is_1yo):
            if age_col not in df.columns or rp_col not in df.columns:
                return 0, 0, 0

            vax_mask = df[rp_col].astype(str).str.lower().str.contains(month_pattern, na=False)

            def row_filter(age_val):
                if "Penta" in vax_name or "OPV" in vax_name:
                    m, d = parse_penta_opv_age(age_val)
                    if m is None:
                        return False
                    if vax_name in ["Penta3", "OPV3"]:
                        if is_1yo:
                            return (m > 3 or (m == 3 and d >= 12)) and m < 12
                        return (m > 3 or (m == 3 and d >= 12)) and m < 60
                    if vax_name in ["Penta1", "OPV1"]:
                        return (m > 1 or (m == 1 and d >= 12)) and m < 60
                    return False
                else:
                    try:
                        numeric_age = float(str(age_val).strip())
                        if pd.isna(numeric_age):
                            return False
                        if vax_name == "Measles1":
                            if is_1yo:
                                return 9 <= numeric_age < 12
                            return 9 <= numeric_age < 60
                        if vax_name == "Measles2":
                            return 10 <= numeric_age < 60
                    except (ValueError, TypeError):
                        return False
                return False

            age_mask = df[age_col].apply(row_filter)
            filtered = df[vax_mask & age_mask]

            m_cnt = int((filtered["GENDER"].astype(str).str.lower().str.startswith("m", na=False)).sum())
            f_cnt = int((filtered["GENDER"].astype(str).str.lower().str.startswith("f", na=False)).sum())
            t_cnt = int(len(filtered))
            return m_cnt, f_cnt, t_cnt

        ind_list = [
            ("R1.5 Number of 1yo (Penta3)", "Penta3_Age", "Penta3_REPORTING_PERIOD", "Penta3", True),
            ("R1.6 Number of 1yo (Measles1)", "MMR1_Age", "MMR1_REPORTING_PERIOD", "Measles1", True),
            ("IR1.1.5.1 Number of U5 (Penta1)", "Penta1_Age", "Penta1_REPORTING_PERIOD", "Penta1", False),
            ("IR1.1.5.2 Number of U5 (Penta3)", "Penta3_Age", "Penta3_REPORTING_PERIOD", "Penta3", False),
            ("IR1.1.5.4 Number of U5 (Measles1)", "MMR1_Age", "MMR1_REPORTING_PERIOD", "Measles1", False),
            ("IR1.1.5.4.1 Number of U5 (Measles2)", "MMR2_Age", "MMR2_REPORTING_PERIOD", "Measles2", False),
        ]

        rows = []
        for label, age_c, rp_c, v_name, is_1yo in ind_list:
            m, f, t = count_vax_gender(age_c, rp_c, v_name, is_1yo)
            rows.append({"Indicator": label, "Male": int(m), "Female": int(f), "Count": int(t)})

        fd_month_blank = df["Full Dose Reporting Month"].isna() | df["Full Dose Reporting Month"].astype(str).str.strip().isin(["", "nan", "None"])
        mmr2_in_period = df["MMR2_REPORTING_PERIOD"].astype(str).str.strip().isin(target_periods) if "MMR2_REPORTING_PERIOD" in df.columns else pd.Series([False]*len(df))

        is_fd_u1_status = df["REACH Full Dose <1"].astype(str).str.strip() == "Full Dose" if "REACH Full Dose <1" in df.columns else pd.Series([False]*len(df))
        fd_u1_month_match = df["Full Dose Reporting Month"].astype(str).str.strip().isin(target_periods) if "Full Dose Reporting Month" in df.columns else pd.Series([False]*len(df))
        fd_under1_df = df[is_fd_u1_status & (fd_u1_month_match | (fd_month_blank & mmr2_in_period))]

        m_fd_u1 = int((fd_under1_df["GENDER"].astype(str).str.lower().str.startswith("m", na=False)).sum())
        f_fd_u1 = int((fd_under1_df["GENDER"].astype(str).str.lower().str.startswith("f", na=False)).sum())
        rows.append({
            "Indicator": "IR1.1.5.5a Number of Full Dose (<1 Year)",
            "Male": m_fd_u1, "Female": f_fd_u1, "Count": int(len(fd_under1_df))
        })

        is_fd_1to5_status = df["REACH Full Dose 1 To 5"].astype(str).str.strip() == "Full Dose" if "REACH Full Dose 1 To 5" in df.columns else pd.Series([False]*len(df))
        fd_1to5_month_match = df["Full Dose Reporting Month"].astype(str).str.strip().isin(target_periods) if "Full Dose Reporting Month" in df.columns else pd.Series([False]*len(df))
        fd_1to5_df = df[is_fd_1to5_status & (fd_1to5_month_match | (fd_month_blank & mmr2_in_period))]

        m_fd_1to5 = int((fd_1to5_df["GENDER"].astype(str).str.lower().str.startswith("m", na=False)).sum())
        f_fd_1to5 = int((fd_1to5_df["GENDER"].astype(str).str.lower().str.startswith("f", na=False)).sum())
        rows.append({
            "Indicator": "IR1.1.5.5b Number of Full Dose (1 To 5 Years)",
            "Male": m_fd_1to5, "Female": f_fd_1to5, "Count": int(len(fd_1to5_df))
        })

        if filter_type == "Quarterly":
            q_code = q_sel.split()[0]
            target_cohort = f"{q_code} {target_year}"
            cohort_mask = df["New_Comer_YN"] == target_cohort if "New_Comer_YN" in df.columns else pd.Series(False, index=df.index)
        else:
            q_code = "Q1" if month_list[0] in ["jan", "feb", "mar"] else "Q2" if month_list[0] in ["apr", "may", "jun"] else "Q3" if month_list[0] in ["jul", "aug", "sep"] else "Q4"
            target_cohort = f"{q_code} {target_year}"
            cohort_mask = (df["New_Comer_YN"] == target_cohort) if "New_Comer_YN" in df.columns else pd.Series(False, index=df.index)

        alo_col = "At Least One Dose" if "At Least One Dose" in df.columns else None
        age_col = "At Least One Dose Month" if "At Least One Dose Month" in df.columns else None

        if alo_col and age_col and "At Least One Dose Date" in df.columns:
            alo_yes = df[alo_col].astype(str).str.strip().str.lower() == "yes"
            alo_dates = pd.to_datetime(df["At Least One Dose Date"], format="%d/%m/%Y", errors="coerce")
            month_match = alo_dates.dt.strftime("%b").str.lower().isin(month_list) & (alo_dates.dt.year == target_year)
            age_num = pd.to_numeric(df[age_col], errors="coerce")
            age_ok = (age_num >= 0) & (age_num < 60)
            alo_df = df[cohort_mask & alo_yes & month_match & age_ok]
        elif "New_Comer_YN" in df.columns:
            alo_df = df[cohort_mask]
        else:
            alo_df = pd.DataFrame()

        m_alo = int((alo_df["GENDER"].astype(str).str.lower().str.startswith("m", na=False)).sum())
        f_alo = int((alo_df["GENDER"].astype(str).str.lower().str.startswith("f", na=False)).sum())
        rows.append({
            "Indicator": "IR1.1.5.6 Number of U5 At Least One Dose",
            "Male": m_alo, "Female": f_alo, "Count": len(alo_df)
        })

        res_df = pd.DataFrame(rows)

        p1_rows = res_df.loc[res_df["Indicator"].str.contains("IR1.1.5.1"), "Count"]
        p3_rows = res_df.loc[res_df["Indicator"].str.contains("IR1.1.5.2"), "Count"]
        p1_val = p1_rows.values[0] if not p1_rows.empty else 0
        p3_val = p3_rows.values[0] if not p3_rows.empty else 0
        dropout = round(((p1_val - p3_val) / p1_val * 100), 2) if p1_val > 0 else 0
        dropout_row = pd.DataFrame([{"Indicator": "IR1.1.5.3 Penta Dropout (%)", "Male": "-", "Female": "-", "Count": dropout}])
        
        return pd.concat([res_df, dropout_row], ignore_index=True)

    st.header(f"🏢 MFM Results for {display_label} {target_year}")
    final_key = "mfm_processed_v_final"

    if final_key in st.session_state and st.session_state[final_key] is not None:
        data = st.session_state[final_key].copy()
        
        global_townships = set(data["TOWNSHIP"].dropna().astype(str).unique()) if "TOWNSHIP" in data.columns else set()
        ts_options = ["All Townships"] + sorted(list(global_townships))
        selected_township = st.selectbox("📍 Select Township Filter:", options=ts_options, key="hei_township_filter")

        filtered_df = data.copy()
        if selected_township != "All Townships" and "TOWNSHIP" in filtered_df.columns:
            filtered_df = filtered_df[filtered_df["TOWNSHIP"] == selected_township]

        if not filtered_df.empty:
            res_df = calculate_indicators(filtered_df, target_months)
            st.dataframe(res_df, use_container_width=True, hide_index=True)
        else:
            st.warning(f"No records found for '{selected_township}'.")
    else:
        st.info("Please process and validate MFM data first.")